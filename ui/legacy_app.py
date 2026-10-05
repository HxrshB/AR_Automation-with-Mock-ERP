from __future__ import annotations

import io                                                             
import json
from pathlib import Path
import sys
import tempfile

import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from app.epicor.mock_client import MockEpicorClient
from app.extractor import detect_document_type, extract_fields
from app.ocr import ocr_pdf
from app.validator import validate


st.set_page_config(
    page_title="AR Invoice Validation",
    page_icon="🧾",
    layout="wide",
    initial_sidebar_state="expanded",
)


# -----------------------------------------------------------------------------
# Session state
# -----------------------------------------------------------------------------
if "mock_client" not in st.session_state:
    st.session_state.mock_client = MockEpicorClient(ROOT / "data" / "mock_epicor_queue.json")
if "logged_in" not in st.session_state:
    st.session_state.logged_in = False
if "username" not in st.session_state:
    st.session_state.username = ""
if "queue_results" not in st.session_state:
    st.session_state.queue_results = {}
if "workflow_status" not in st.session_state:
    st.session_state.workflow_status = {}
if "selected_ticket" not in st.session_state:
    st.session_state.selected_ticket = None
if "last_rejection_reason" not in st.session_state:
    st.session_state.last_rejection_reason = ""


def load_rates() -> dict[str, float]:
    path = ROOT / "data" / "mock_erp.json"
    data = json.loads(path.read_text(encoding="utf-8"))
    return {str(k): float(v) for k, v in data.get("discount_rates", {}).items()}


def money(value):
    return f"SAR {value:,.2f}" if value is not None else "—"


def value_or_dash(value):
    return value if value not in (None, "") else "—"


def save_upload(uploaded, path: Path):
    path.write_bytes(uploaded.getbuffer())


def process_documents(accounting_path: Path, jst_path: Path):
    accounting_pages = ocr_pdf(accounting_path)
    jst_pages = ocr_pdf(jst_path)

    if detect_document_type(accounting_pages) != "accounting":
        raise ValueError("The Accounting PDF could not be confirmed as an Accounting Information document.")
    if detect_document_type(jst_pages) != "jst":
        raise ValueError("The JST PDF could not be confirmed as a JST / Service Report document.")

    accounting = extract_fields(accounting_pages, "accounting")
    jst = extract_fields(jst_pages, "jst")
    result = validate(accounting, jst, load_rates())
    return result


def process_ticket(ticket_id: str):
    if ticket_id in st.session_state.queue_results:
        return st.session_state.queue_results[ticket_id]

    client = st.session_state.mock_client
    attachments = client.get_attachments(ticket_id)
    accounting_attachment = None
    jst_attachment = None

    for attachment in attachments:
        doc_type = (attachment.document_type or "").lower()
        name = attachment.name.lower()
        if doc_type == "accounting" or "account" in name:
            accounting_attachment = attachment
        elif doc_type == "jst" or "jst" in name or "service" in name:
            jst_attachment = attachment

    if not accounting_attachment or not jst_attachment:
        return {
            "ticket_id": ticket_id,
            "validation": None,
            "error": "Required Accounting and/or JST attachment is missing.",
        }

    try:
        validation = process_documents(accounting_attachment.path, jst_attachment.path)
        result = {
            "ticket_id": ticket_id,
            "validation": validation,
            "error": None,
            "accounting_file": accounting_attachment.name,
            "jst_file": jst_attachment.name,
        }
    except Exception as exc:
        result = {
            "ticket_id": ticket_id,
            "validation": None,
            "error": str(exc),
            "accounting_file": accounting_attachment.name,
            "jst_file": jst_attachment.name,
        }

    st.session_state.queue_results[ticket_id] = result
    return result


def get_queue():
    if not st.session_state.logged_in:
        return []
    return st.session_state.mock_client.get_user_queue(st.session_state.username)


def workflow_status(ticket_id: str) -> str:
    return st.session_state.workflow_status.get(ticket_id, "PENDING")


def set_workflow_status(ticket_id: str, status: str):
    st.session_state.workflow_status[ticket_id] = status


def validation_status(result) -> str:
    if result is None:
        return "REVIEW"
    if result.get("error"):
        return "REVIEW"
    return result["validation"].status


def validation_badge(status: str):
    if status == "VALID":
        st.success("✓ VALID")
    elif status == "INVALID":
        st.error("✗ INVALID")
    else:
        st.warning("⚠ REVIEW")


def render_check(name: str, value):
    if value is True:
        st.success(f"✓ {name}")
    elif value is False:
        st.error(f"✗ {name}")
    else:
        st.warning(f"⚠ {name} — Unable to verify")


def render_validation_detail(ticket_id: str):
    result = process_ticket(ticket_id)
    if result.get("error"):
        st.error(result["error"])
        return

    validation = result["validation"]
    st.caption(f"Ticket: {ticket_id}")

    validation_badge(validation.status)

    checks = [
        ("Request Number", validation.request_match),
        ("DRSS Number", validation.drss_match),
        ("Job ID", validation.job_id_match),
        ("Well Name", validation.well_match),
        ("Rig Short Name", validation.rig_short_name_match),
        ("Contract Number", validation.contract_match),
        ("Total Amount", validation.total_amount_match),
        ("Discount Rate / Amount", validation.discount_percent_match),
        ("Final Total Calculation", validation.total_after_discount_match),
    ]

    st.subheader("Validation Checks")
    cols = st.columns(3)
    for i, (name, value) in enumerate(checks):
        with cols[i % 3]:
            render_check(name, value)

    st.divider()
    st.subheader("Document Comparison")
    a, b = st.columns(2)
    with a:
        st.markdown("### Accounting")
        st.write(f"**Request:** {value_or_dash(validation.accounting.request_number)}")
        st.write(f"**Well:** {value_or_dash(validation.accounting.well_name)}")
        st.write(f"**Rig:** {value_or_dash(validation.accounting.rig_short_name)}")
        st.write(f"**Contract:** {value_or_dash(validation.accounting.contract_no)}")
        st.write(f"**Total:** {money(validation.accounting.grand_total_after_discount)}")
    with b:
        st.markdown("### JST / Service Report")
        st.write(f"**DRSS:** {value_or_dash(validation.jst.drss_number)}")
        st.write(f"**Well:** {value_or_dash(validation.jst.well_name)}")
        st.write(f"**Rig:** {value_or_dash(validation.jst.rig_short_name)}")
        st.write(f"**Contract:** {value_or_dash(validation.jst.contract_no)}")
        st.write(f"**Job ID:** {value_or_dash(validation.jst.job_number)}")
        st.write(f"**Total:** {money(validation.jst.grand_total_after_discount)}")

    st.divider()
    st.subheader("Financial Verification")
    rate = validation.reference.get("expected_discount_percent")
    before = validation.jst.grand_total_before_discount
    expected = None if rate is None or before is None else round(before * rate / 100, 2)
    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Contract", value_or_dash(validation.reference.get("contract_no")))
    c2.metric("Configured Discount", f"{rate:.2f}%" if rate is not None else "Not configured")
    c3.metric("Expected Discount", money(expected))
    c4.metric("JST Discount", money(validation.jst.discount_amount))

    c1, c2, c3 = st.columns(3)
    c1.metric("Before Discount", money(validation.jst.grand_total_before_discount))
    c2.metric("Accounting Total", money(validation.accounting.grand_total_after_discount))
    c3.metric("JST Total", money(validation.jst.grand_total_after_discount))

    if validation.warnings:
        st.subheader("Warnings")
        for warning in validation.warnings:
            st.warning(warning)
    if validation.reasons:
        st.subheader("Review / Failure Reasons")
        for reason in validation.reasons:
            st.write(f"• {reason}")

    st.caption(
        f"Job dates: {value_or_dash(validation.jst.start_date)} → {value_or_dash(validation.jst.final_date)}. "
        "Dates are informational and do not affect validation status."
    )

    with st.expander("Extracted Accounting Information"):
        st.json(validation.accounting.model_dump(exclude={"raw_text_pages"}))
    with st.expander("Extracted JST Information"):
        st.json(validation.jst.model_dump(exclude={"raw_text_pages"}))
    with st.expander("OCR Text — Accounting"):
        for page, text_value in validation.accounting.raw_text_pages.items():
            st.markdown(f"**Page {page}**")
            st.text(text_value)
    with st.expander("OCR Text — JST"):
        for page, text_value in validation.jst.raw_text_pages.items():
            st.markdown(f"**Page {page}**")
            st.text(text_value)

    st.divider()
    st.subheader("Workflow Action")
    current = workflow_status(ticket_id)
    if current == "APPROVED":
        st.success("Ticket approved in the mock Epicor workflow.")
        return
    if current == "REJECTED":
        st.error("Ticket rejected in the mock Epicor workflow.")
        if st.session_state.last_rejection_reason:
            st.caption(f"Reason: {st.session_state.last_rejection_reason}")
        return

    action_cols = st.columns(3)
    with action_cols[0]:
        if st.button("✓ Approve", key=f"approve_{ticket_id}", type="primary", use_container_width=True):
            if validation.status != "VALID":
                st.warning("Only VALID invoices can be automatically approved in this prototype.")
            else:
                st.session_state.mock_client.approve_ticket(ticket_id, "Approved by validation UI")
                set_workflow_status(ticket_id, "APPROVED")
                st.rerun()
    with action_cols[1]:
        if st.button("⚠ Mark Review", key=f"review_{ticket_id}", use_container_width=True):
            set_workflow_status(ticket_id, "REVIEW")
            st.rerun()
    with action_cols[2]:
        if st.button("✕ Reject", key=f"reject_{ticket_id}", use_container_width=True):
            st.session_state[f"show_reject_{ticket_id}"] = True

    if st.session_state.get(f"show_reject_{ticket_id}"):
        reason = st.text_area(
            "Rejection reason",
            value="; ".join(validation.reasons) if validation.reasons else "Validation failed.",
            key=f"reject_reason_{ticket_id}",
        )
        c1, c2 = st.columns(2)
        with c1:
            if st.button("Confirm Rejection", key=f"confirm_reject_{ticket_id}", type="primary", use_container_width=True):
                if not reason.strip():
                    st.error("A rejection reason is required.")
                else:
                    st.session_state.mock_client.reject_ticket(ticket_id, reason)
                    st.session_state.last_rejection_reason = reason
                    set_workflow_status(ticket_id, "REJECTED")
                    st.session_state[f"show_reject_{ticket_id}"] = False
                    st.rerun()
        with c2:
            if st.button("Cancel", key=f"cancel_reject_{ticket_id}", use_container_width=True):
                st.session_state[f"show_reject_{ticket_id}"] = False
                st.rerun()


def build_excel(queue_rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Validation Queue"
    headers = [
        "Ticket ID", "Request / DRSS", "Well", "Rig", "Contract", "Start Date",
        "Final Date", "Job ID", "Discount %", "Expected Discount", "JST Discount",
        "Accounting Total", "JST Total", "Validation", "Workflow Status", "Warnings",
        "Reason", "Action"
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
        cell.alignment = Alignment(horizontal="center")

    for row in queue_rows:
        ws.append([row.get(h, "") for h in headers])

    widths = [20, 16, 16, 14, 16, 14, 14, 12, 12, 18, 16, 18, 14, 14, 16, 36, 55, 14]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + idx) if idx <= 26 else "A"].width = width
    ws.freeze_panes = "A2"

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return out.getvalue()


def queue_rows(tickets):
    rows = []
    for ticket in tickets:
        result = process_ticket(ticket.ticket_id)
        validation = result.get("validation")
        if validation:
            rate = validation.reference.get("expected_discount_percent")
            before = validation.jst.grand_total_before_discount
            expected = None if rate is None or before is None else round(before * rate / 100, 2)
            rows.append({
                "Ticket ID": ticket.ticket_id,
                "Request / DRSS": validation.accounting.request_number or validation.jst.request_number or "",
                "Well": validation.jst.well_name or validation.accounting.well_name or "",
                "Rig": validation.jst.rig_short_name or validation.accounting.rig_short_name or "",
                "Contract": validation.accounting.contract_no or validation.jst.contract_no or "",
                "Start Date": validation.jst.start_date or validation.accounting.start_date or "",
                "Final Date": validation.jst.final_date or validation.accounting.final_date or "",
                "Job ID": validation.jst.job_number or "",
                "Discount %": rate if rate is not None else "",
                "Expected Discount": expected if expected is not None else "",
                "JST Discount": validation.jst.discount_amount if validation.jst.discount_amount is not None else "",
                "Accounting Total": validation.accounting.grand_total_after_discount if validation.accounting.grand_total_after_discount is not None else "",
                "JST Total": validation.jst.grand_total_after_discount if validation.jst.grand_total_after_discount is not None else "",
                "Validation": validation.status,
                "Workflow Status": workflow_status(ticket.ticket_id),
                "Warnings": "; ".join(validation.warnings),
                "Reason": "; ".join(validation.reasons),
                "Action": "Approve" if validation.status == "VALID" else "Review",
            })
        else:
            rows.append({
                "Ticket ID": ticket.ticket_id,
                "Validation": "REVIEW",
                "Workflow Status": workflow_status(ticket.ticket_id),
                "Reason": result.get("error", ""),
                "Action": "Review",
            })
    return rows


# -----------------------------------------------------------------------------
# Sidebar / login
# -----------------------------------------------------------------------------
with st.sidebar:
    st.title("🧾 AR Automation")
    st.caption("Invoice validation prototype")
    st.divider()

    if not st.session_state.logged_in:
        st.subheader("Epicor Login")
        username = st.text_input("Username", value="demo_user")
        password = st.text_input("Password", type="password")
        if st.button("Connect", type="primary", use_container_width=True):
            if st.session_state.mock_client.login(username, password):
                st.session_state.logged_in = True
                st.session_state.username = username
                st.rerun()
        st.info("Prototype mode: this uses the local Mock Epicor client. Real Epicor authentication will replace this later.")
        st.stop()

    st.success(f"Connected as {st.session_state.username}")
    if st.button("Disconnect", use_container_width=True):
        st.session_state.logged_in = False
        st.session_state.username = ""
        st.session_state.queue_results = {}
        st.session_state.selected_ticket = None
        st.rerun()

    page = st.radio(
        "Navigation",
        ["Dashboard", "Validation Queue", "Manual Validation", "Reports", "Discount Master"],
        label_visibility="collapsed",
    )

# -----------------------------------------------------------------------------
# Dashboard
# -----------------------------------------------------------------------------
if page == "Dashboard":
    st.title("AR Invoice Validation Dashboard")
    st.caption("Unified interface for the existing OCR/validation engine and the mock queue.")

    tickets = get_queue()
    # Include already actioned tickets that remain in session state by processing
    # the original mock queue data when available.
    all_mock = st.session_state.mock_client._tickets.values()
    for ticket in all_mock:
        if ticket.ticket_id not in [t.ticket_id for t in tickets] and ticket.ticket_id in st.session_state.workflow_status:
            tickets.append(ticket)

    rows = queue_rows(tickets)
    validation_counts = {"VALID": 0, "REVIEW": 0, "INVALID": 0}
    workflow_counts = {"PENDING": 0, "APPROVED": 0, "REJECTED": 0, "REVIEW": 0}
    for row in rows:
        validation_counts[row.get("Validation", "REVIEW")] = validation_counts.get(row.get("Validation", "REVIEW"), 0) + 1
        workflow_counts[row.get("Workflow Status", "PENDING")] = workflow_counts.get(row.get("Workflow Status", "PENDING"), 0) + 1

    c1, c2, c3, c4 = st.columns(4)
    c1.metric("Queue Items", len(rows))
    c2.metric("Valid", validation_counts["VALID"])
    c3.metric("Review", validation_counts["REVIEW"])
    c4.metric("Invalid", validation_counts["INVALID"])

    st.subheader("Workflow")
    c1, c2, c3 = st.columns(3)
    c1.metric("Pending", workflow_counts["PENDING"])
    c2.metric("Approved", workflow_counts["APPROVED"])
    c3.metric("Rejected", workflow_counts["REJECTED"])

    if rows:
        st.subheader("Current Queue")
        st.dataframe(rows, use_container_width=True, hide_index=True)
    else:
        st.info("No queue items are currently available for this user.")

# -----------------------------------------------------------------------------
# Validation Queue
# -----------------------------------------------------------------------------
elif page == "Validation Queue":
    st.title("Validation Queue")
    st.caption("Review documents, inspect validation evidence, and perform mock workflow actions.")

    tickets = list(st.session_state.mock_client._tickets.values())
    rows = queue_rows(tickets)

    f1, f2, f3 = st.columns([2, 2, 3])
    with f1:
        status_filter = st.selectbox("Validation Status", ["All", "VALID", "REVIEW", "INVALID"])
    with f2:
        workflow_filter = st.selectbox("Workflow Status", ["All", "PENDING", "REVIEW", "APPROVED", "REJECTED"])
    with f3:
        search = st.text_input("Search Ticket / Request / Well / Rig / Contract")

    filtered = []
    for row in rows:
        if status_filter != "All" and row.get("Validation") != status_filter:
            continue
        if workflow_filter != "All" and row.get("Workflow Status") != workflow_filter:
            continue
        if search:
            haystack = " ".join(str(v) for v in row.values()).lower()
            if search.lower() not in haystack:
                continue
        filtered.append(row)

    st.write(f"Showing **{len(filtered)}** item(s)")
    st.dataframe(filtered, use_container_width=True, hide_index=True)

    if filtered:
        options = [r["Ticket ID"] for r in filtered]
        selected = st.selectbox("Open invoice", options, index=0)
        st.session_state.selected_ticket = selected
        if st.button("Open Validation Details", type="primary"):
            st.session_state.selected_ticket = selected

        if st.session_state.selected_ticket:
            st.divider()
            render_validation_detail(st.session_state.selected_ticket)

# -----------------------------------------------------------------------------
# Manual Validation
# -----------------------------------------------------------------------------
elif page == "Manual Validation":
    st.title("Manual Invoice Validation")
    st.caption("Upload the Accounting Information PDF and the matching JST / Service Report PDF.")

    accounting_file = st.file_uploader("Accounting Information PDF", type=["pdf"], key="manual_accounting")
    jst_file = st.file_uploader("JST / Service Report PDF", type=["pdf"], key="manual_jst")

    if st.button("Validate Documents", type="primary", use_container_width=True):
        if not accounting_file or not jst_file:
            st.error("Please upload both PDF documents.")
            st.stop()

        with tempfile.TemporaryDirectory() as temp_dir:
            temp_dir = Path(temp_dir)
            accounting_path = temp_dir / "accounting.pdf"
            jst_path = temp_dir / "jst.pdf"
            save_upload(accounting_file, accounting_path)
            save_upload(jst_file, jst_path)
            try:
                with st.spinner("Reading and validating documents..."):
                    result = process_documents(accounting_path, jst_path)
                st.session_state.manual_result = result
            except Exception as exc:
                st.error(f"An error occurred while processing the documents: {exc}")

    if "manual_result" in st.session_state:
        result = st.session_state.manual_result
        st.divider()
        st.subheader("Validation Result")
        validation_badge(result.status)
        checks = [
            ("Request Number", result.request_match),
            ("DRSS Number", result.drss_match),
            ("Job ID", result.job_id_match),
            ("Well Name", result.well_match),
            ("Rig Short Name", result.rig_short_name_match),
            ("Contract Number", result.contract_match),
            ("Total Amount", result.total_amount_match),
            ("Discount Rate / Amount", result.discount_percent_match),
            ("Final Total Calculation", result.total_after_discount_match),
        ]
        cols = st.columns(3)
        for i, (name, value) in enumerate(checks):
            with cols[i % 3]:
                render_check(name, value)

        st.subheader("Financial Verification")
        rate = result.reference.get("expected_discount_percent")
        before = result.jst.grand_total_before_discount
        expected = None if rate is None or before is None else round(before * rate / 100, 2)
        c1, c2, c3, c4 = st.columns(4)
        c1.metric("Contract", value_or_dash(result.reference.get("contract_no")))
        c2.metric("Configured Discount", f"{rate:.2f}%" if rate is not None else "Not configured")
        c3.metric("Expected Discount", money(expected))
        c4.metric("JST Discount", money(result.jst.discount_amount))

        if result.warnings:
            for warning in result.warnings:
                st.warning(warning)
        if result.reasons:
            for reason in result.reasons:
                st.write(f"• {reason}")

        with st.expander("Extracted Accounting Information"):
            st.json(result.accounting.model_dump(exclude={"raw_text_pages"}))
        with st.expander("Extracted JST Information"):
            st.json(result.jst.model_dump(exclude={"raw_text_pages"}))

# -----------------------------------------------------------------------------
# Reports
# -----------------------------------------------------------------------------
elif page == "Reports":
    st.title("Reports & Excel Export")
    st.caption("Export the same queue information used by the UI for AR review/reporting.")
    tickets = list(st.session_state.mock_client._tickets.values())
    rows = queue_rows(tickets)
    if not rows:
        st.info("No queue records to export.")
    else:
        st.dataframe(rows, use_container_width=True, hide_index=True)
        xlsx = build_excel(rows)
        st.download_button(
            "Download Validation Queue Excel",
            data=xlsx,
            file_name="validation_queue.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            type="primary",
            use_container_width=True,
        )

# -----------------------------------------------------------------------------
# Discount Master
# -----------------------------------------------------------------------------
elif page == "Discount Master":
    st.title("Discount Master")
    st.caption("Current prototype contract → discount configuration used by the deterministic validator.")
    rates = load_rates()
    rows = [{"Contract No.": contract, "Discount %": rate} for contract, rate in rates.items()]
    st.dataframe(rows, use_container_width=True, hide_index=True)
    st.info("Editing is intentionally read-only in this first UI integration. We will add controlled master-data editing after the queue workflow is stable.")
